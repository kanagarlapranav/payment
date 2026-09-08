import openpyxl
from openpyxl.styles import Font, PatternFill
from database.queries import get_all_transactions
import os

def generate_excel_report(output_path: str):
    """Generates an Excel report of all transactions."""
    transactions = get_all_transactions()
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Transactions"
    
    # Headers
    headers = [
        "ID", "Date", "Time", "Type", "Person", "Sender", "Recipient", 
        "Amount", "UPI ID", "Phone", "Reference Number", "Transaction ID", 
        "Payment App", "Bank", "Status", "Balance Before", "Balance After", "Created At"
    ]
    
    ws.append(headers)
    
    # Styling headers
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
    
    # Freeze top row
    ws.freeze_panes = 'A2'
    
    # Data
    for t in transactions:
        row = [
            t['id'], 
            t['transaction_date'], 
            t['transaction_time'], 
            t['transaction_type'],
            t['person_name'], 
            t['sender_name'], 
            t['recipient_name'], 
            t['amount'], 
            t['upi_id'], 
            t['phone_number'],
            t['reference_number'], 
            t['transaction_id'], 
            t['payment_app'], 
            t['bank_name'],
            t['payment_status'], 
            t['balance_before'], 
            t['balance_after'], 
            t['created_at']
        ]
        ws.append(row)
        
    # Formatting
    for row in range(2, len(transactions) + 2):
        # Format amount, balances as currency
        ws.cell(row=row, column=8).number_format = '₹#,##0.00'
        ws.cell(row=row, column=16).number_format = '₹#,##0.00'
        ws.cell(row=row, column=17).number_format = '₹#,##0.00'
        
    # Auto-adjust column widths
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter # Get the column name
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column].width = adjusted_width

    wb.save(output_path)


def generate_pdf_statement(output_path: str):
    """Generates a professional PDF statement of all transactions."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
    from reportlab.pdfgen import canvas
    from database.queries import get_all_transactions, get_balance_setting
    from utils.currency import format_currency
    from utils.dates import get_current_time_in_tz, format_display_date

    transactions = get_all_transactions()
    current_balance = get_balance_setting()

    total_sent = sum(t['amount'] for t in transactions if t['transaction_type'] == 'SENT')
    total_received = sum(t['amount'] for t in transactions if t['transaction_type'] == 'RECEIVED')
    net_flow = total_received - total_sent
    now_str = get_current_time_in_tz().strftime("%d %b %Y, %I:%M %p")

    class StatementCanvas(canvas.Canvas):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self._saved_page_states = []

        def showPage(self):
            self._saved_page_states.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            num_pages = len(self._saved_page_states)
            for state in self._saved_page_states:
                self.__dict__.update(state)
                self.draw_decorations(num_pages)
                super().showPage()
            super().save()

        def draw_decorations(self, total_pages):
            self.saveState()
            self.setFont("Helvetica", 8)
            self.setFillColor(colors.HexColor("#718096"))
            self.setStrokeColor(colors.HexColor("#E2E8F0"))
            self.setLineWidth(0.5)
            self.line(36, 40, 559, 40)
            self.drawString(36, 28, "Payment Tracker • Automated Financial Statement")
            self.drawRightString(559, 28, f"Page {self._pageNumber} of {total_pages}")
            self.restoreState()

    doc = SimpleDocTemplate(
        output_path,
        pagesize=A4,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=50
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'TitleStyle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#1A202C")
    )
    subtitle_style = ParagraphStyle(
        'SubTitleStyle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#718096")
    )
    card_label = ParagraphStyle(
        'CardLabel',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        leading=10,
        textColor=colors.HexColor("#718096")
    )
    card_val = ParagraphStyle(
        'CardVal',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=14,
        textColor=colors.HexColor("#1A202C")
    )
    th_style = ParagraphStyle(
        'TH',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.white
    )
    td_style = ParagraphStyle(
        'TD',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        leading=10,
        textColor=colors.HexColor("#2D3748")
    )
    td_bold = ParagraphStyle(
        'TDBold',
        parent=td_style,
        fontName='Helvetica-Bold'
    )
    td_sent = ParagraphStyle(
        'TDSent',
        parent=td_bold,
        textColor=colors.HexColor("#E53E3E")
    )
    td_recv = ParagraphStyle(
        'TDRecv',
        parent=td_bold,
        textColor=colors.HexColor("#38A169")
    )

    story = []
    story.append(Paragraph("Account Statement", title_style))
    story.append(Paragraph(f"Generated on {now_str} • Personal Payment Tracker", subtitle_style))
    story.append(Spacer(1, 8))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#3182CE"), spaceBefore=2, spaceAfter=8))

    # Summary Cards Box
    summary_data = [
        [
            Paragraph("CURRENT BALANCE", card_label),
            Paragraph("TOTAL SENT", card_label),
            Paragraph("TOTAL RECEIVED", card_label),
            Paragraph("NET SAVINGS", card_label),
            Paragraph("TRANSACTIONS", card_label),
        ],
        [
            Paragraph(format_currency(current_balance), card_val),
            Paragraph(f"<font color='#E53E3E'>{format_currency(total_sent)}</font>", card_val),
            Paragraph(f"<font color='#38A169'>{format_currency(total_received)}</font>", card_val),
            Paragraph(format_currency(net_flow), card_val),
            Paragraph(str(len(transactions)), card_val),
        ]
    ]
    summary_table = Table(summary_data, colWidths=[104, 104, 104, 104, 107])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#F7FAFC")),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor("#E2E8F0")),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#E2E8F0")),
        ('TOPPADDING', (0,0), (-1,-1), 6),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ('LEFTPADDING', (0,0), (-1,-1), 6),
        ('RIGHTPADDING', (0,0), (-1,-1), 6),
    ]))
    story.append(summary_table)
    story.append(Spacer(1, 12))

    # Table of transactions
    th_row = [
        Paragraph("#", th_style),
        Paragraph("Date & Time", th_style),
        Paragraph("Type", th_style),
        Paragraph("Person / Entity", th_style),
        Paragraph("App / Bank", th_style),
        Paragraph("Reference", th_style),
        Paragraph("Amount", th_style),
        Paragraph("Balance After", th_style),
    ]
    table_rows = [th_row]

    for idx, t in enumerate(transactions, 1):
        is_sent = t['transaction_type'] == 'SENT'
        type_style = td_sent if is_sent else td_recv
        amt_prefix = "- " if is_sent else "+ "
        amt_text = f"{amt_prefix}{format_currency(t['amount'])}"
        
        date_display = format_display_date(t['transaction_date'])
        time_display = t['transaction_time'] or ""
        dt_text = f"{date_display}<br/>{time_display}" if time_display else date_display
        
        app_bank = t['payment_app'] or ""
        if t['bank_name']:
            app_bank = f"{app_bank} • {t['bank_name']}" if app_bank else t['bank_name']
            
        ref_text = t['reference_number'] or t['transaction_id'] or "—"
        
        row = [
            Paragraph(str(idx), td_style),
            Paragraph(dt_text, td_style),
            Paragraph(t['transaction_type'], type_style),
            Paragraph(t['person_name'] or "Unknown", td_bold),
            Paragraph(app_bank or "—", td_style),
            Paragraph(ref_text[:14], td_style),
            Paragraph(amt_text, type_style),
            Paragraph(format_currency(t['balance_after']), td_style),
        ]
        table_rows.append(row)

    tx_table = Table(table_rows, colWidths=[20, 75, 45, 110, 85, 70, 58, 60], repeatRows=1)
    tx_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#2B6CB0")),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E0")),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ('LEFTPADDING', (0,0), (-1,-1), 4),
        ('RIGHTPADDING', (0,0), (-1,-1), 4),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor("#F7FAFC")]),
    ]))
    story.append(tx_table)

    doc.build(story, canvasmaker=StatementCanvas)
